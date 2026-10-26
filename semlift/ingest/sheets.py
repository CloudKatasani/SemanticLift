"""USAGE, MAP and SRC-CAT adapters (XLSX or CSV, header located in the first 10 rows)."""
from __future__ import annotations

import csv
import datetime as dt
import re
from pathlib import Path
from typing import Any, Iterator

from ..model import MetaGraph, Package, Provenance, SourceColumn, SourceObject, SourceSystem

DATE_FORMATS = ("%b %d, %Y %I:%M:%S %p", "%b %d, %Y", "%Y-%m-%d %H:%M:%S", "%Y-%m-%d", "%m/%d/%Y")


def parse_date(v: Any) -> str | None:
    if v is None or v == "":
        return None
    if isinstance(v, dt.datetime):
        return v.date().isoformat()
    if isinstance(v, dt.date):
        return v.isoformat()
    s = str(v).strip()
    for f in DATE_FORMATS:
        try:
            return dt.datetime.strptime(s, f).date().isoformat()
        except ValueError:
            continue
    return None


def num(v: Any) -> float | None:
    if v is None or v == "":
        return None
    try:
        return float(str(v).replace(",", ""))
    except ValueError:
        return None


def iter_rows(path: Path) -> Iterator[tuple[str, int, list]]:
    """Yield (sheet, 1-based row number, values) for every row of every sheet/CSV."""
    if path.suffix.lower() in (".xlsx", ".xlsm"):
        import openpyxl
        wb = openpyxl.load_workbook(path, read_only=True, data_only=True)
        for ws in wb.worksheets:
            for i, row in enumerate(ws.iter_rows(values_only=True), start=1):
                yield ws.title, i, list(row)
        wb.close()
    else:
        with open(path, newline="", encoding="utf-8-sig") as fh:
            for i, row in enumerate(csv.reader(fh), start=1):
                yield path.stem, i, row


def table_rows(path: Path, required: set[str]) -> Iterator[tuple[str, dict, list]]:
    """Locate the header row (first 10 rows of a sheet) containing ``required`` and yield dict rows.
    Also yields title rows above the header as ('__title__', ...)."""
    header: list[str] | None = None
    cur_sheet = None
    for sheet, i, row in iter_rows(path):
        if sheet != cur_sheet:
            cur_sheet, header = sheet, None
        cells = ["" if v is None else str(v).strip() for v in row]
        if header is None:
            if i <= 10 and required <= {c.lower() for c in cells}:
                header = [c.lower() for c in cells]
            elif i <= 10 and any(cells):
                yield "__title__", {"text": " ".join(c for c in cells if c), "loc": f"{sheet}!{i}"}, row
            continue
        if not any(cells):
            continue
        yield f"{sheet}!{i}", dict(zip(header, row)), row


def parse_usage(path: Path, rel: str, g: MetaGraph) -> int:
    n = 0
    for loc, rec, _ in table_rows(path, {"report_name", "runs_per_day"}):
        if loc == "__title__":
            t = rec["text"]
            m = re.search(r"\((\d+)\s+reports?\)", t, re.I)
            if m:
                g.stated_counts["usage_title_reports"] = int(m.group(1))
            m = re.search(r"(\d+)\s+Distinct", t, re.I)
            if m:
                g.stated_counts["usage_title_per_category"] = int(m.group(1))
            continue
        name = str(rec.get("report_name") or "").strip()
        if not name:
            continue
        g.usage_rows.append({
            "report_name": name, "report_path": str(rec.get("report_path") or "").strip(),
            "package": rec.get("package"), "report_type": rec.get("report_type"),
            "folder_category": rec.get("folder_category"), "duplicate_group_id": rec.get("duplicate_group_id"),
            "similarity_score": num(rec.get("similarity_score")), "query_name": rec.get("query_name"),
            "column_count": num(rec.get("column_count")), "runs_per_day": num(rec.get("runs_per_day")),
            "last_run_date": parse_date(rec.get("last_run_date")), "avg_duration_min": num(rec.get("avg_duration_min")),
            "loc": f"{rel}:{loc}",
        })
        n += 1
    return n


VENDOR_ALIASES = {"oracle": "Oracle", "sap hana": "SAP HANA", "osisoft pi": "OSIsoft PI", "salesforce": "Salesforce",
                  "sql server": "SQL Server", "db2": "Db2", "teradata": "Teradata"}


def parse_mapping(path: Path, rel: str, g: MetaGraph) -> int:
    n = 0
    for loc, rec, _ in table_rows(path, {"fm package", "source database/system"}):
        if loc == "__title__":
            continue
        unit = str(rec.get("fm package/data module") or rec.get("fm package") or "").strip()
        pkg = str(rec.get("fm package") or unit).strip()
        src = str(rec.get("source database/system") or "").strip()
        if not pkg or not src:
            continue
        vendor, _, db = src.partition(" - ")
        vendor = VENDOR_ALIASES.get(vendor.strip().lower(), vendor.strip())
        db = db.strip() or vendor
        prov = Provenance(rel, loc)
        g.source_systems.setdefault(db, SourceSystem(name=db, vendor=vendor, database=db, origin="MAP", prov=prov))
        p = g.packages.setdefault(pkg, Package(name=pkg, prov=prov))
        p.source_system = p.source_system or db
        if unit != pkg and unit not in p.data_modules:
            p.data_modules.append(unit)
        n += 1
    return n


def parse_catalog(path: Path, rel: str, g: MetaGraph) -> int:
    """SRC-CAT CSV: table_schema/table_catalog, table_name, column_name, data_type, [row_count]."""
    n = 0
    for loc, rec, _ in table_rows(path, {"table_name", "column_name"}):
        if loc == "__title__":
            continue
        schema = str(rec.get("table_schema") or rec.get("schema") or "").strip().upper()
        db = str(rec.get("table_catalog") or rec.get("database") or schema).strip().upper()
        table = str(rec.get("table_name") or "").strip().upper()
        col = str(rec.get("column_name") or "").strip().upper()
        if not table or not col:
            continue
        key = f"CAT:{schema}.{table}"
        so = g.source_objects.get(key)
        if so is None:
            so = SourceObject(key=key, system=db, schema=schema, table=table, origin="CATALOG", prov=Provenance(rel, loc))
            g.source_objects[key] = so
        rc = num(rec.get("row_count"))
        if rc is not None:
            so.row_count = int(rc)
        if so.column(col) is None:
            so.columns.append(SourceColumn(col, str(rec.get("data_type") or "").strip()))
        n += 1
    g.catalog_loaded = g.catalog_loaded or n > 0
    return n
