"""design_workbook.xlsx for client sign-off (FR-18)."""
from __future__ import annotations

import datetime as dt
from pathlib import Path

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

from .. import __version__
from ..design.model import Design
from ..model import MetaGraph, TargetObject

HEAD = PatternFill("solid", fgColor="1F3A5F")
HFONT = Font(bold=True, color="FFFFFF")


def _sheet(wb: Workbook, title: str, headers: list[str], rows: list[list]) -> None:
    ws = wb.create_sheet(title[:31])
    ws.append(headers)
    for c in ws[1]:
        c.fill, c.font = HEAD, HFONT
        c.alignment = Alignment(vertical="top")
    for r in rows:
        ws.append(["" if v is None else (", ".join(map(str, v)) if isinstance(v, (list, tuple, set)) else v) for v in r])
    ws.freeze_panes = "A2"
    ws.auto_filter.ref = ws.dimensions
    for i, h in enumerate(headers, start=1):
        width = max([len(str(h))] + [min(len(str(r[i - 1])) if i - 1 < len(r) and r[i - 1] is not None else 0, 80) for r in rows[:300]])
        ws.column_dimensions[get_column_letter(i)].width = min(max(10, width + 2), 80)


def write_workbook(path: Path, g: MetaGraph, d: Design, objects: list[TargetObject], trace: list[dict],
                   decisions: dict | None = None) -> None:
    decisions = decisions or {}
    wb = Workbook()
    ws = wb.active
    ws.title = "Summary"
    rat = g.analysis.get("rationalization", {})
    summary = [
        ("Tool", f"SemLift {__version__}"), ("Reports", len(g.reports)), ("Intent clusters", rat.get("clusters")),
        ("Retired clusters", rat.get("retired_clusters")), ("Logical entities", len(g.entities)),
        ("FM-defined entities", sum(e.origin == "FM" for e in g.entities.values())),
        ("Inferred entities", sum(e.origin != "FM" for e in g.entities.values())),
        ("CORE objects", len(d.tables)), ("Aggregates", len(d.aggregates)), ("Semantic views", len(d.semantic_views)),
        ("Policies", len(d.policies)), ("Data products", len(d.products)), ("Target objects", len(objects)),
        ("Stub objects (not auto-deployable)", sum(o.status == "stub" for o in objects)),
        ("Findings", len(g.findings)), ("Blocking findings", sum(f.blocking for f in g.findings)),
    ]
    ws.append(["Metric", "Value"])
    for c in ws[1]:
        c.fill, c.font = HEAD, HFONT
    for k, v in summary:
        ws.append([k, v])
    ws.column_dimensions["A"].width = 40
    ws.column_dimensions["B"].width = 30

    _sheet(wb, "Object Inventory",
           ["Order", "Object", "Layer", "Type", "Domain", "Status", "Review", "Confidence", "Owner", "Classification",
            "Rules", "Depends on", "Comment", "File"],
           [[o.order, o.id, o.layer, o.object_type, o.domain, o.status, decisions.get(o.id, {}).get("status", ""),
             o.confidence, o.owner, o.classification, o.rule_ids, o.depends_on, o.comment, o.meta.get("file")]
            for o in sorted(objects, key=lambda o: o.order)])
    cm = []
    for tid, t in sorted(d.tables.items()):
        for c in t.columns:
            cm.append([tid, c.name, c.kind, c.data_type, c.ref, c.source_column, c.expr, c.masking_policy,
                       ", ".join(f"{k}={v}" for k, v in sorted(c.tags.items())), "yes" if c.todo else "", c.confidence, c.comment])
    _sheet(wb, "Column Mapping", ["Table", "Column", "Kind", "Type", "Cognos item", "Source column", "Expression",
                                  "Masking policy", "Tags", "TODO", "Confidence", "Comment"], cm)
    _sheet(wb, "Measure Catalog",
           ["Entity", "Measure", "Origin", "FM regularAggregate", "Default aggregation", "Additivity", "Unit", "Rule",
            "Calculated", "Ratio parts", "Confidence", "Expression"],
           [[m.entity, m.name, m.origin, m.regular_aggregate, m.default_aggregation, m.additivity, m.unit, m.rule,
             m.is_calculated, m.ratio_parts, m.confidence, m.expression] for m in sorted(g.measures.values(), key=lambda m: (m.entity, m.name))])
    _sheet(wb, "Aggregate Plan",
           ["Aggregate", "Fact", "Grain", "Implementation", "Coverage %", "Clusters", "Est. rows", "Credits/day", "Target lag",
            "Excluded grain", "Status"],
           [[a.id, a.fact, a.grain_refs, a.implementation, a.coverage_pct, a.clusters, a.est_rows, a.credits_per_day,
             a.target_lag, a.excluded, "stub" if a.stub else "proposed"] for a in sorted(d.aggregates.values(), key=lambda a: a.id)])
    fs = g.analysis.get("aggregate_facts", {})
    _sheet(wb, "Fact Coverage", ["Fact", "Clusters", "Weighted demand", "Covered", "Coverage %", "Aggregates", "Excluded (fan trap / no path)"],
           [[f, s["clusters"], s["total_weight"], s["covered_weight"], s["coverage_pct"], s["aggregates"], s["excluded_attributes"]]
            for f, s in sorted(fs.items())])
    sec = []
    for r in sorted(g.security_rules.values(), key=lambda r: (r.type, r.name)):
        pol = [p.id for p in d.policies.values() if p.cognos_rule == r.name]
        if r.type == "column_grant":
            pol = pol or sorted({c.masking_policy for t in d.tables.values() for c in t.columns
                                 if c.tags.get("ACCESS_GROUP") and c.ref == f"{r.target_entity}.{r.target_item}"})
        if r.type == "abac":
            pol = ["ABAC clauses in MP_* (tag ABAC_POLICY)"]
        if r.type == "namespace":
            pol = ["SEC.GRANTS"]
        sec.append([r.type, r.name, f"{r.target_entity}.{r.target_item}".strip("."), r.predicate, r.roles_in, r.roles_exempt,
                    r.roles_denied, r.masked_value, r.origin, pol,
                    r.prov.source_file if r.prov else ""])
    _sheet(wb, "Security Matrix", ["Type", "Cognos rule", "Target", "Predicate / condition", "Roles in", "Exempt", "Denied",
                                   "Masked value", "Origin", "Platform policy", "Source file"], sec)
    _sheet(wb, "Report Disposition",
           ["Report", "Package", "Type", "Folder", "Cluster", "Canonical", "Disposition", "Runs/day", "Last run", "Avg min",
            "Similarity", "Defects"],
           [[r.name, r.package, r.type, r.folder, r.cluster, r.canonical, r.disposition, r.runs_per_day, r.last_run,
             r.avg_duration, r.similarity, r.defects] for r in sorted(g.reports.values(), key=lambda r: (r.package, r.cluster, r.name))])
    port = g.analysis.get("portfolio")
    if port:
        _sheet(wb, "Rationalization",
               ["Report", "Package", "Folder", "Cluster", "Runs/day", "Group runs/day", "Last run", "Recommendation", "Rule",
                "Decision", "Disposition", "Consolidate into", "KPIs", "Reasons", "Also note"],
               [[r["report"], r["package"], r["folder"], r["cluster"], r["runs_per_day"], r.get("group_runs_per_day"), r["last_run"],
                 r["recommendation"], r["rule"], (r.get("decision") or {}).get("disposition", ""), r["disposition"],
                 r.get("target") or "", r["kpis"], " | ".join(r["reasons"]), " | ".join(r["flags"])] for r in port["reports"]])
    cat = g.analysis.get("kpi_catalog")
    if cat:
        rows = []
        for gr in cat["groups"]:
            for v in gr["variants"]:
                rows.append([gr["title"], gr["kind"], gr["status"], v["kpi"], v["formula"], v["aggregation"], v["filter"],
                             len(v["reports"]), v["live_reports"], "yes" if v["recommended"] else "",
                             "yes" if v["id"] == gr.get("approved_variant") else ""])
        _sheet(wb, "Common KPIs", ["KPI group", "Kind", "Status", "KPI item", "Definition", "Aggregation", "Filter rule",
                                   "Reports", "Live reports", "Recommended", "Approved"], rows)
    src = g.analysis.get("source_inventory")
    if src:
        _sheet(wb, "Data Sources", ["Database", "Vendor", "Schema", "Table / view", "Type", "Basis", "Columns", "Keep reports",
                                    "Review reports", "Review only", "Scope", "Landed as", "CORE objects"],
               [[t["database"], t["vendor"], t["schema"], t["table"], t["object_type"], t["basis"], [c["column"] for c in t["columns"]],
                 len(t["keep_reports"]), len(t["review_reports"]), t["review_only"], t["scope"], t["landed"], t["core"]]
                for t in src["tables"]])
    _sheet(wb, "Intent Clusters",
           ["Cluster", "Package", "Canonical", "Members", "Disposition", "Runs/day", "Weighted demand", "Facts", "Items", "Reason"],
           [[c.id, c.package, c.canonical, c.members, c.disposition, c.runs_per_day, c.weighted_demand, c.facts, len(c.items), c.reason]
            for c in sorted(g.clusters.values(), key=lambda c: c.id)])
    _sheet(wb, "Findings", ["ID", "Code", "Severity", "Title", "Object", "Evidence", "Owner", "Resolution", "Status", "Count", "Blocking"],
           [[f.id, f.code, f.severity, f.title, f.object, f.evidence, f.owner, f.resolution, f.status, f.count, f.blocking]
            for f in g.findings])
    _sheet(wb, "Numbers That Change", ["Report", "Measure", "Cognos", "Semantic", "Origin"],
           [[r["report"], r["measure"], r["cognos"], r["semantic"], r["origin"]] for r in g.analysis.get("numbers_that_change", [])])
    _sheet(wb, "Validation", ["Rule", "Check", "Passed", "Detail", "Failures"],
           [[v["rule"], v["title"], v["passed"], v["detail"], v["failures"][:20]] for v in g.analysis.get("validation", [])])
    if trace:
        keys = list(trace[0].keys())
        _sheet(wb, "Traceability", keys, [[r[k] for k in keys] for r in trace])
    fixed = dt.datetime(2000, 1, 1)
    wb.properties.created = fixed
    wb.properties.modified = fixed
    wb.properties.creator = f"SemLift {__version__}"
    wb.save(path)
